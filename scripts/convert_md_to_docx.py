#!/usr/bin/env python3
"""
scripts/convert_md_to_docx.py
모든 콘텐츠 .md 파일을 .docx 형식으로 변환한다.
출력 위치: {project_root}/docx/
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = Path(__file__).parent.parent   # project root
DOCX_DIR = BASE_DIR / "docx"

EXCLUDE_NAMES = {
    "CLAUDE.md", "README.md", "SESSION_PLAN.md",
    "RESEARCH_GUIDE.md", "TRANSLATION_GUIDE.md",
    "TEMPLATE_EXCHANGE.md", "TEMPLATE_SYSTEM.md", "TEMPLATE_CLEARING.md",
    "INVENTORY.md",
}

# (src_dir relative to BASE_DIR, dst_dir relative to DOCX_DIR, label)
DOMAIN_MAPPINGS = [
    ("md/01_equities/en",  "01_equities/en",  "01_equities EN"),
    ("md/01_equities/ko",  "01_equities/ko",  "01_equities KO"),
    ("md/02_derivatives/en", "02_derivatives/en", "02_derivatives EN"),
    ("md/02_derivatives/ko", "02_derivatives/ko", "02_derivatives KO"),
    ("md/03_clearing/en",  "03_clearing/en",  "03_clearing EN"),
    ("md/03_clearing/ko",  "03_clearing/ko",  "03_clearing KO"),
    ("md/_common",         "_common/en",      "_common EN"),
    ("md/_common/ko",      "_common/ko",      "_common KO"),
]


def parse_table(lines):
    """Parse markdown table lines into rows of cells."""
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.split("|")[1:-1]]
            rows.append(cells)
    # Remove separator row (contains ---)
    rows = [r for r in rows if not all(re.match(r'^[-:]+$', c) for c in r)]
    return rows


def add_table_to_doc(doc, rows):
    """Add a table to the document."""
    if not rows or len(rows) < 1:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.cell(i, j)
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    paragraph.style = doc.styles['Normal']
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                # Bold header row
                if i == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True


def md_to_docx(md_path, docx_path):
    """Convert a markdown file to docx."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)

    lines = content.split('\n')
    i = 0
    table_buffer = []
    in_table = False
    in_code_block = False
    code_buffer = []

    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                code_text = '\n'.join(code_buffer)
                p = doc.add_paragraph()
                p.style = doc.styles['Normal']
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_buffer = []
                in_code_block = False
            else:
                # Flush table if any
                if in_table and table_buffer:
                    rows = parse_table(table_buffer)
                    add_table_to_doc(doc, rows)
                    table_buffer = []
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        # Table handling
        if line.strip().startswith('|'):
            in_table = True
            table_buffer.append(line)
            i += 1
            continue
        elif in_table:
            # End of table
            rows = parse_table(table_buffer)
            add_table_to_doc(doc, rows)
            doc.add_paragraph()  # spacing
            table_buffer = []
            in_table = False
            # Don't increment i, process current line

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # Blockquote (treat as italic paragraph)
        if line.startswith('>'):
            text = line.lstrip('> ').strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.italic = True
            run.font.size = Pt(9)
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ('---', '***', '___'):
            doc.add_paragraph('─' * 60)
            i += 1
            continue

        # Bullet points
        if re.match(r'^(\s*)[-*]\s', line):
            indent_level = len(line) - len(line.lstrip())
            text = re.sub(r'^(\s*)[-*]\s+', '', line).strip()
            # Clean markdown bold/italic
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
            if indent_level >= 4:
                p.paragraph_format.left_indent = Inches(0.5)
            i += 1
            continue

        # Numbered list
        if re.match(r'^(\s*)\d+\.\s', line):
            text = re.sub(r'^(\s*)\d+\.\s+', '', line).strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            doc.add_paragraph(text, style='List Number')
            i += 1
            continue

        # Regular paragraph
        text = line.strip()
        # Clean markdown formatting
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        text = re.sub(r'\*(.+?)\*', r'\1', text)
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)  # links
        if text:
            doc.add_paragraph(text)
        i += 1

    # Flush remaining table
    if in_table and table_buffer:
        rows = parse_table(table_buffer)
        add_table_to_doc(doc, rows)

    doc.save(str(docx_path))


def process_domain(src_rel: str, dst_rel: str, label: str) -> int:
    """Convert all .md files in src_dir (relative to BASE_DIR) to docx in dst_dir."""
    src_dir = BASE_DIR / src_rel
    dst_dir = DOCX_DIR / dst_rel

    if not src_dir.exists():
        print(f"  [{label}] Source not found: {src_dir}, skipping.")
        return 0

    dst_dir.mkdir(parents=True, exist_ok=True)

    md_files = sorted(src_dir.rglob('*.md'))
    count = 0
    for md_file in md_files:
        if md_file.name in EXCLUDE_NAMES:
            continue
        # Skip files inside docx subdirectories
        if 'docx' in md_file.parts:
            continue
        # Skip _common/ko files when processing _common (EN) to avoid duplication
        if src_rel == "md/_common":
            rel_parts = list(md_file.relative_to(src_dir).parts)
            if rel_parts and rel_parts[0] == "ko":
                continue

        rel = md_file.relative_to(src_dir)
        parts = list(rel.parts)
        if len(parts) > 1:
            docx_name = '_'.join(parts[:-1]) + '_' + parts[-1].replace('.md', '.docx')
        else:
            docx_name = parts[0].replace('.md', '.docx')

        docx_path = dst_dir / docx_name
        try:
            md_to_docx(md_file, docx_path)
            count += 1
            print(f"  [{label}] {rel} -> {docx_name}")
        except Exception as e:
            print(f"  [{label}] ERROR {rel}: {e}")
    return count


def main():
    total = 0
    for src_rel, dst_rel, label in DOMAIN_MAPPINGS:
        print(f"\nConverting {label}...")
        count = process_domain(src_rel, dst_rel, label)
        total += count
        print(f"  -> {count} files")
    print(f"\nDone! Total: {total} files")


if __name__ == '__main__':
    main()
